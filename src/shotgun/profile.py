"""The candidate profile: parse a resume into structured YAML, then load it.

The profile is the single source of truth for tailoring. It lives in
private/profile.yaml, which is gitignored — it has your address and phone
number in it.
"""

from __future__ import annotations

import logging
from pathlib import Path

from pydantic import BaseModel, Field

from . import llm
from .config import PRIVATE_DIR

log = logging.getLogger(__name__)

PROFILE_PATH = PRIVATE_DIR / "profile.yaml"


# --------------------------------------------------------------- schema

class Contact(BaseModel):
    name: str
    email: str
    phone: str | None = None
    location: str | None = Field(default=None, description="City, Country")
    linkedin: str | None = None
    github: str | None = None
    website: str | None = None


class Role(BaseModel):
    company: str
    title: str
    start: str = Field(description="YYYY-MM or YYYY")
    end: str | None = Field(default=None, description="YYYY-MM, YYYY, or null if current")
    location: str | None = None
    bullets: list[str] = Field(description="Achievements as written on the source resume")


class Education(BaseModel):
    institution: str
    qualification: str
    year: str | None = None


class Profile(BaseModel):
    contact: Contact
    headline: str | None = None
    summary: str | None = None
    roles: list[Role] = Field(default_factory=list)
    skills: list[str] = Field(default_factory=list)
    certifications: list[str] = Field(default_factory=list)
    education: list[Education] = Field(default_factory=list)
    # Was present in profile.yaml but missing from this model, so pydantic
    # dropped it on load and save() erased it from the file. Language
    # proficiency is load-bearing for DACH and EU roles, and plenty of
    # application forms ask outright.
    languages: list[str] = Field(
        default_factory=list,
        description="One per line, e.g. 'German — B1', 'English — Professional'",
    )
    # Not parsed from the resume — you fill these in by hand.
    work_authorization: dict[str, str] = Field(
        default_factory=dict,
        description="Country code -> status, e.g. {IN: citizen, GB: needs sponsorship}",
    )
    salary_expectation: dict[str, str] = Field(default_factory=dict)
    notice_period: str | None = None


# ----------------------------------------------------------- extraction

EXTRACT_SYSTEM = """Extract the resume text into the given schema.

Rules:
- Copy achievements verbatim into `bullets`. Do not rewrite, embellish, or \
invent anything. Tailoring happens later, from this faithful record.
- If a field is genuinely absent from the resume, leave it null or empty. Never \
guess at a phone number, a date, or an employer.
- Preserve the original ordering of roles, most recent first."""


def read_text(path: Path) -> str:
    """Pull plain text out of a PDF, DOCX, or text file."""
    suffix = path.suffix.lower()

    if suffix == ".pdf":
        import pdfplumber

        with pdfplumber.open(path) as pdf:
            pages = [page.extract_text() or "" for page in pdf.pages]
        return "\n\n".join(pages)

    if suffix in (".docx", ".doc"):
        import docx

        document = docx.Document(str(path))
        parts = [p.text for p in document.paragraphs]
        for table in document.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        return "\n".join(parts)

    return path.read_text(errors="replace")


def parse_resume(path: Path) -> Profile:
    """Resume file -> structured Profile, via Claude."""
    text = read_text(path)
    if len(text.strip()) < 200:
        raise ValueError(
            f"Only extracted {len(text.strip())} characters from {path.name}. "
            "If it's a scanned PDF there is no text layer to read — export a "
            "text-based PDF or point at a DOCX."
        )

    # No profile to cache here — this is what produces it — so the first
    # system block is the instructions.
    return llm.parse_structured(
        [EXTRACT_SYSTEM],
        f"<resume>\n{text}\n</resume>",
        Profile,
        max_tokens=16000,
        label="parse_resume",
    )


# ------------------------------------------------------------ persistence

def save(profile: Profile, path: Path = PROFILE_PATH) -> Path:
    import yaml

    path.parent.mkdir(parents=True, exist_ok=True)
    path.write_text(
        yaml.safe_dump(profile.model_dump(), sort_keys=False, allow_unicode=True)
    )
    return path


# The fields no resume contains, so no re-parse can recover them.
MANUAL_FIELDS = ("work_authorization", "salary_expectation", "notice_period")


def carry_manual_fields(
    parsed: Profile, path: Path = PROFILE_PATH
) -> tuple[list[str], Path | None]:
    """Move the hand-filled fields off an existing profile onto a fresh parse.

    `profile init` writes the whole file, so re-parsing an updated resume used
    to silently wipe work authorization, salary expectation and notice period
    — the three things a resume never states and the three the scorer and the
    form filler most depend on. `authorised_countries()` reads
    work_authorization, and with it empty `rank --authorised-only` matches
    nothing at all.

    A backup is written too, because carrying the fields still cannot carry
    the prose. `load_yaml` hands the raw file to the model, comments included,
    so the notes written into this file — why an EU Blue Card makes other
    EU states reachable without employer sponsorship, which resume
    contradiction was resolved and how — are part of the scoring prompt, and
    `yaml.safe_dump` does not keep them. Returns what was carried and where
    the old file went.
    """
    if not path.exists():
        return [], None

    backup = path.with_suffix(path.suffix + ".bak")
    backup.write_text(path.read_text())

    try:
        existing = load(path)
    except Exception:
        # A profile too broken to load is still worth having backed up, and
        # the fresh parse is no worse than what is there.
        return [], backup

    carried = []
    for field in MANUAL_FIELDS:
        value = getattr(existing, field, None)
        if not value:
            continue
        setattr(parsed, field, value)
        carried.append(field)
    return carried, backup


def _unmap_colon_strings(value):
    """Repair list entries YAML turned into single-key mappings.

    An unquoted colon in a hand-edited list item is a YAML mapping, not a
    string: `- Microsoft Certified: Azure Fundamentals (AZ-900)` parses as
    `{"Microsoft Certified": "Azure Fundamentals (AZ-900)"}`. The README tells
    you to open this file and edit it by hand, so this is a matter of when not
    if, and the raw pydantic error names only `certifications.8`.
    """
    if isinstance(value, dict) and len(value) == 1:
        (key, inner), = value.items()
        if isinstance(inner, str):
            return f"{key}: {inner}"
    return value


def load(path: Path = PROFILE_PATH) -> Profile:
    import yaml

    if not path.exists():
        raise FileNotFoundError(
            f"No profile at {path}. Run `shotgun profile init <resume.pdf>` first."
        )

    raw = yaml.safe_load(path.read_text()) or {}

    for field in ("skills", "certifications"):
        if isinstance(raw.get(field), list):
            raw[field] = [_unmap_colon_strings(v) for v in raw[field]]
    for role in raw.get("roles") or []:
        if isinstance(role, dict) and isinstance(role.get("bullets"), list):
            role["bullets"] = [_unmap_colon_strings(b) for b in role["bullets"]]

    return Profile.model_validate(raw)


def load_yaml(path: Path = PROFILE_PATH) -> str:
    """The raw YAML, for the cached system block.

    Returned verbatim so the bytes are identical on every call in a run — a
    re-serialize here would risk key reordering and blow the prompt cache.
    """
    if not path.exists():
        raise FileNotFoundError(
            f"No profile at {path}. Run `shotgun profile init <resume.pdf>` first."
        )
    return path.read_text()


def template() -> Profile:
    """An empty profile to fill in by hand."""
    return Profile(
        contact=Contact(name="", email=""),
        roles=[Role(company="", title="", start="", bullets=[""])],
        education=[Education(institution="", qualification="")],
    )
